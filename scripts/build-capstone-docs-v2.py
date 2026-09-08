"""Create the optional printable assignment directly from the maintained JSON."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
d=json.loads((ROOT/'CAPSTONE/source/bus311-capstone.json').read_text())
doc=Document();sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 1'].font.size=Pt(20);doc.styles['Heading 2'].font.size=Pt(13)
def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title):
 if len(doc.paragraphs)>0: doc.add_page_break()
 doc.add_heading(title,1)
def footer():
 f=sec.footer.paragraphs[0];f.text='BUS311 Company Capstone · Fall 2026 · '
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
footer()
doc.add_heading('BUS311 Company Capstone',0)
p(d['hub']['purpose'])
p('Individual project · 100 points · 25% of the course grade')
p(d['project']['companyEligibility'])
h('One continuing workbook, then one presentation')
p('Use the project workbook for Stages 1–3. It contains Start, Revenue, Model, Sources, Revision, and an annotated Example. Use the six-slide template for the final presentation. This printable brief is optional.')
for s in d['hub']['stages']:
 h(s['label']+': '+s['title']+' ('+str(s['points'])+' points)')
 p(s['dateLabel']);p(s['work']);p('Submit: '+s['submit'])
p('Project page: '+d['hub']['publicUrl'])
page('A manageable analysis')
h('Research minimum')
p('Three comparable years of revenue, two drivers, one peer comparison, two risks, and one forecast claim. Connect each to the proposed decision. Additional research is optional and earns credit only by improving the published criteria.')
for x in d['decisionPaths']:
 h(x['title']);p(x['question']);p(x['method']);p(x['inputs']);p(x['boundary'])
h('How the model works')
p('Compare the change with doing nothing. Enter the initial cash flow, annual benefits and costs, tax and discount rates, and end-of-project cash flows. The workbook calculates five-year NPV. The status quo has zero incremental NPV. Compare base, upside, and downside assumptions and test the path-appropriate sensitivity. Explain the model’s simplifying assumptions and your source IDs.')
p(d['feedback']['support'])
h('One challenge')
p(d['policies']['challenge'])
for t in d['challengePrompts']:p(t)
page('Grading and feedback')
p('Workbook milestones: 25 points. PowerPoint: 50 points. Oral Presentation: 25 points.')
for c in d['rubric']['criteria']:
 h(c['title']+' — '+str(c['points'])+' points');p(c['description'])
p(d['rubric']['partialCredit'])
h('Time to use feedback')
p('Model due November 4. Feedback target: '+d['feedback']['target']+'. '+d['feedback']['revision'])
page('The six-slide decision briefing')
for i,s in enumerate(d['slideOutline'],1):
 h(str(i)+'. '+s['title']);p(s['prompt']);p('Evidence: '+s['evidence'])
p('Use six core slides plus an optional supporting appendix. Up to seven minutes presenting and three minutes answering questions. The template notes suggest timing.')
p('Upload '+d['finalSubmission']['files'][0]+' by '+d['finalSubmission']['deadlineLabel']+'. Present the same file on your assigned date.')
p(d['feedback']['finalModel'])
page('Sources and submission policies')
for key in ['individualWork','stageOne','sources','factSetPermission','aiPrivacy','missingEvidence','lateWork','revision','finalFileLock','absence','submission']:p(d['policies'][key])
p('FactSet setup: '+d['hub']['factSetSetupUrl'])
doc.save(ROOT/'CAPSTONE/bus311-capstone-assignment.docx')
print('Built optional five-page assignment brief.')
